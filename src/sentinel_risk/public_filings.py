"""Reproducible, original-filing credit case from cached SEC XBRL companyfacts.

All financial values are USD, not millions. No scorecard, rating or PD is inferred.
The case uses annual filings available at a fixed historical cutoff, not live data.
"""
from __future__ import annotations

import argparse
import csv
from datetime import date, datetime, timezone
import gzip
import hashlib
import io
import json
from pathlib import Path
from statistics import median
import time
import urllib.request

ROOT = Path(__file__).resolve().parents[2]
DATA = ROOT / "data" / "public_filings"
AS_OF = "2025-04-01"
YEARS = (2022, 2023, 2024)
ISSUERS = {"KO": "0000021344", "PEP": "0000077476", "KDP": "0001418135"}

# Priority is explicit; concepts are never added together indiscriminately.
FIELDS = {
    "revenue": ("flow", ["RevenueFromContractWithCustomerExcludingAssessedTax", "Revenues", "SalesRevenueNet"]),
    "gross_profit": ("flow", ["GrossProfit"]),
    "operating_income": ("flow", ["OperatingIncomeLoss"]),
    "net_income": ("flow", ["NetIncomeLoss"]),
    "cost_of_goods_sold": ("flow", ["CostOfGoodsAndServicesSold", "CostOfRevenue"]),
    "total_assets": ("instant", ["Assets"]),
    "current_assets": ("instant", ["AssetsCurrent"]),
    "current_liabilities": ("instant", ["LiabilitiesCurrent"]),
    "total_liabilities": ("instant", ["Liabilities"]),
    "equity": ("instant", ["StockholdersEquity"]),
    "consolidated_equity": ("instant", ["StockholdersEquityIncludingPortionAttributableToNoncontrollingInterest"]),
    "cash": ("instant", ["CashAndCashEquivalentsAtCarryingValue"]),
    "inventory": ("instant", ["InventoryNet"]),
    "accounts_receivable": ("instant", ["AccountsReceivableNetCurrent", "AccountsNotesAndLoansReceivableNetCurrent"]),
    "operating_cash_flow": ("flow", ["NetCashProvidedByUsedInOperatingActivities"]),
    "capital_expenditure": ("flow", ["PaymentsToAcquirePropertyPlantAndEquipment", "PaymentsToAcquireProductiveAssets"]),
    "interest_expense": ("flow", ["InterestExpense", "InterestExpenseNonoperating"]),
}
RATIOS = {
    "gross_margin": ("Gross margin", "gross_profit / revenue", "%"),
    "operating_margin": ("Operating margin", "operating_income / revenue", "%"),
    "net_margin": ("Net margin", "net_income / revenue", "%"),
    "current_ratio": ("Current ratio", "current_assets / current_liabilities", "x"),
    "cash_ratio": ("Cash ratio", "cash / current_liabilities", "x"),
    "quick_ratio": ("Quick ratio (CA less inventories)", "(current_assets - inventory) / current_liabilities", "x"),
    "liabilities_to_assets": ("Liabilities / assets", "total_liabilities / total_assets", "%"),
    "cfo_margin": ("Operating cash flow margin", "operating_cash_flow / revenue", "%"),
    "fcf_margin": ("Free cash flow margin", "(operating_cash_flow - capital_expenditure) / revenue", "%"),
    "cfo_to_liabilities": ("Operating cash flow / liabilities", "operating_cash_flow / total_liabilities", "%"),
    "revenue_growth": ("Revenue growth", "revenue / prior_year_revenue - 1", "%"),
}
CONTEXT = {
    "KO": {
        "profile": "Global beverage company with concentrate and finished-product operations. Bottling arrangements and refranchising affect comparisons with integrated peers.",
        "event": "The 2024 filing attributes the operating cash flow decline primarily to a $6.0 billion IRS tax litigation deposit, alongside currency, tax and refranchising effects. Reported cash flow is retained without adding the deposit back.",
        "source": "https://investors.coca-colacompany.com/filings-reports/annual-filings-10-k/content/0000021344-25-000011/ko-20241231.htm",
        "section": "Business; MD&A — Liquidity, Capital Resources and Financial Position",
    },
    "PEP": {
        "profile": "Global convenient foods and beverages company. Its foods operations make it an imperfect beverage peer; fiscal 2024 ended December 28 rather than December 31.",
        "event": "Review the foods and beverages mix, geographic performance and demand alongside consolidated margins. The reported interest expense concept is not used for a cross-company gross interest coverage comparison.",
        "source": "https://www.sec.gov/Archives/edgar/data/77476/000007747625000007/pep-20241228.htm",
        "section": "Business; MD&A; Consolidated Financial Statements",
    },
    "KDP": {
        "profile": "Beverage and coffee company reporting U.S. Refreshment Beverages, U.S. Coffee and International segments. Coffee mix and geography differ from the other two issuers.",
        "event": "The 2024 results describe $718 million of goodwill and intangible impairments and a $225 million accrual for GHOST distribution termination payments. These help explain lower GAAP operating income; reported GAAP figures are retained.",
        "source": "https://investors.keurigdrpepper.com/2025-02-25-Keurig-Dr-Pepper-Reports-Q4-and-Full-Year-2024-Results-and-Provides-2025-Outlook",
        "section": "2024 Full Year Consolidated Results and Segment Results",
    },
}


def download_sources(data_dir: Path = DATA, user_agent: str = "") -> list[dict]:
    """Refresh only on explicit CLI request; identify a real contact to the SEC."""
    if "@" not in user_agent:
        raise ValueError("Pass --user-agent with your project name and real contact email")
    (data_dir / "raw").mkdir(parents=True, exist_ok=True)
    sources = []
    for ticker, cik in ISSUERS.items():
        url = f"https://data.sec.gov/api/xbrl/companyfacts/CIK{cik}.json"
        request = urllib.request.Request(url, headers={"User-Agent": user_agent})
        with urllib.request.urlopen(request, timeout=60) as response:
            raw = response.read()
        document = json.loads(raw)
        if str(document["cik"]).zfill(10) != cik:
            raise ValueError(f"Unexpected issuer for {ticker}")
        # Content-addressed files keep an existing manifest valid if a later fetch fails.
        digest = hashlib.sha256(raw).hexdigest()
        cache = f"raw/{ticker}-{digest[:12]}.json.gz"
        (data_dir / cache).write_bytes(gzip.compress(raw, mtime=0))
        sources.append({"ticker": ticker, "cik": cik, "entity_name": document["entityName"],
                        "url": url, "retrieved_at": datetime.now(timezone.utc).isoformat(),
                        "sha256": digest, "cache": cache, "bytes": len(raw)})
        time.sleep(0.25)
    temporary = data_dir / "sources.json.tmp"
    temporary.write_text(json.dumps(sources, indent=2) + "\n")
    temporary.replace(data_dir / "sources.json")
    return sources


def load_sources(data_dir: Path = DATA) -> list[tuple[dict, dict]]:
    result = []
    sources = json.loads((data_dir / "sources.json").read_text())
    tickers = [source["ticker"] for source in sources]
    if len(tickers) != len(ISSUERS) or set(tickers) != set(ISSUERS):
        raise ValueError("Source manifest must contain each configured issuer exactly once")
    for source in sources:
        ticker = source["ticker"]
        if source["cik"] != ISSUERS[ticker]:
            raise ValueError(f"Manifest ticker/CIK mismatch: {ticker}")
        expected_url = f"https://data.sec.gov/api/xbrl/companyfacts/CIK{ISSUERS[ticker]}.json"
        if source["url"] != expected_url:
            raise ValueError(f"Unexpected source URL: {ticker}")
        cache = Path(source["cache"])
        resolved = (data_dir / cache).resolve()
        if cache.is_absolute() or not resolved.is_relative_to(data_dir.resolve()):
            raise ValueError(f"Cache path escapes dataset: {ticker}")
        raw = gzip.decompress(resolved.read_bytes())
        if hashlib.sha256(raw).hexdigest() != source["sha256"]:
            raise ValueError(f"Source checksum mismatch: {source['ticker']}")
        document = json.loads(raw)
        if str(document["cik"]).zfill(10) != source["cik"]:
            raise ValueError(f"Source issuer mismatch: {source['ticker']}")
        result.append((source, document))
    return result


def annual_fact(row: dict, year: int, as_of: str) -> bool:
    if row.get("form") != "10-K" or row.get("filed", "9999") > as_of:
        return False
    if not row.get("start") or not row.get("end", "").startswith(str(year)):
        return False
    days = (date.fromisoformat(row["end"]) - date.fromisoformat(row["start"])).days
    return 330 <= days <= 400


def select_annual(document: dict, source: dict, year: int, as_of: str = AS_OF) -> dict:
    """Choose the earliest annual revenue accession, then exact periods within it.

    Do not filter by XBRL fy (it labels the filing year, even for comparatives).
    Never mix a later restatement, a quarter, or a different filing into a row.
    """
    taxonomy = document["facts"]["us-gaap"]
    anchors = [row for tag in FIELDS["revenue"][1]
               for row in taxonomy.get(tag, {}).get("units", {}).get("USD", [])
               if annual_fact(row, year, as_of)]
    if not anchors:
        raise ValueError(f"No annual revenue for {source['ticker']} {year} at {as_of}")
    anchor = min(anchors, key=lambda row: (row["filed"], row["accn"]))
    evidence = {}
    values = {}
    for field, (kind, tags) in FIELDS.items():
        selected = None
        for tag in tags:
            candidates = [row for row in taxonomy.get(tag, {}).get("units", {}).get("USD", [])
                          if row.get("accn") == anchor["accn"] and row.get("end") == anchor["end"]
                          and row.get("filed", "9999") <= as_of
                          and (row.get("start") == anchor["start"] if kind == "flow" else not row.get("start"))]
            if candidates:
                if len({row["val"] for row in candidates}) != 1:
                    raise ValueError(f"Ambiguous {source['ticker']} {year} {tag}")
                selected = candidates[0]
                evidence[field] = {"kind": "reported", "taxonomy": "us-gaap", "concept": tag,
                                   "unit": "USD", **{k: selected[k] for k in ("val", "end", "accn", "filed")},
                                   "start": selected.get("start"), "source_sha256": source["sha256"]}
                break
        values[field] = selected["val"] if selected else None
        if selected is None:
            evidence[field] = {"kind": "unavailable", "reason": "No matching standard-taxonomy USD fact in the selected filing; never filled with zero."}
    if source["ticker"] == "PEP" and values["interest_expense"] is not None:
        # PepsiCo tags the issuer line 'Net interest expense and other' as
        # us-gaap:InterestExpense. It is not comparable with gross interest.
        values["net_interest_and_other"] = values["interest_expense"]
        evidence["net_interest_and_other"] = evidence["interest_expense"] | {
            "issuer_label": "Net interest expense and other"}
        values["interest_expense"] = None
        evidence["interest_expense"] = {"kind": "unavailable", "reason": "The issuer's InterestExpense fact represents net interest expense and other. It is preserved separately as net_interest_and_other; gross interest expense is not sourced here."}
    if values["total_liabilities"] is None and all(values[k] is not None for k in ("total_assets", "consolidated_equity")):
        values["total_liabilities"] = values["total_assets"] - values["consolidated_equity"]
        evidence["total_liabilities"] = {"kind": "derived", "formula": "total_assets - consolidated_equity", "inputs": ["total_assets", "consolidated_equity"]}
    if values["gross_profit"] is None and all(values[k] is not None for k in ("revenue", "cost_of_goods_sold")):
        values["gross_profit"] = values["revenue"] - values["cost_of_goods_sold"]
        evidence["gross_profit"] = {"kind": "derived", "formula": "revenue - cost_of_goods_sold", "inputs": ["revenue", "cost_of_goods_sold"]}
    cik = int(source["cik"])
    accession = anchor["accn"]
    return {"ticker": source["ticker"], "year": year, "period_start": anchor["start"],
            "period_end": anchor["end"], "filed": anchor["filed"], "accession": accession,
            "filing_url": f"https://www.sec.gov/Archives/edgar/data/{cik}/{accession.replace('-', '')}/{accession}-index.html",
            "values": values, "evidence": evidence}


def divide(numerator: float | None, denominator: float | None) -> float | None:
    return numerator / denominator if numerator is not None and denominator is not None and denominator > 0 else None


def difference(left: float | None, right: float | None) -> float | None:
    return left - right if left is not None and right is not None else None


def calculate(row: dict, prior: dict | None = None) -> dict:
    v = row["values"]
    fcf = difference(v["operating_cash_flow"], v["capital_expenditure"])
    growth = divide(v["revenue"], prior["values"]["revenue"]) if prior else None
    return {"gross_margin": divide(v["gross_profit"], v["revenue"]),
            "operating_margin": divide(v["operating_income"], v["revenue"]),
            "net_margin": divide(v["net_income"], v["revenue"]),
            "current_ratio": divide(v["current_assets"], v["current_liabilities"]),
            "cash_ratio": divide(v["cash"], v["current_liabilities"]),
            "quick_ratio": divide(difference(v["current_assets"], v["inventory"]), v["current_liabilities"]),
            "liabilities_to_assets": divide(v["total_liabilities"], v["total_assets"]),
            "cfo_margin": divide(v["operating_cash_flow"], v["revenue"]),
            "fcf_margin": divide(fcf, v["revenue"]),
            "cfo_to_liabilities": divide(v["operating_cash_flow"], v["total_liabilities"]),
            "revenue_growth": growth - 1 if growth is not None else None}


def monitoring(rows: list[dict]) -> list[dict]:
    """Transparent research screens, not contractual covenants or issuer ratings."""
    latest, prior = rows[-1], rows[-2]
    r = latest["ratios"]
    change = difference(r["operating_margin"], prior["ratios"]["operating_margin"])
    cfo_growth = divide(latest["values"]["operating_cash_flow"], prior["values"]["operating_cash_flow"])
    fcf = difference(latest["values"]["operating_cash_flow"], latest["values"]["capital_expenditure"])
    checks = [
        ("Liquidity", "Current ratio below 1.0x", r["current_ratio"], r["current_ratio"] is not None and r["current_ratio"] < 1),
        ("Profitability", "Operating margin falls by more than 2 percentage points", change, change is not None and change < -0.02),
        ("Cash flow", "Operating cash flow falls by more than 20%", cfo_growth - 1 if cfo_growth is not None else None, cfo_growth is not None and cfo_growth < 0.8),
        ("Free cash flow", "Operating cash flow less capital expenditure is negative", fcf, fcf is not None and fcf < 0),
    ]
    return [{"area": area, "rule": rule, "value": value,
             "status": "unavailable" if value is None else "review" if hit else "not_triggered",
             "basis": "Illustrative screening threshold; not a covenant breach"}
            for area, rule, value, hit in checks]


def build_case(data_dir: Path = DATA, as_of: str = AS_OF) -> dict:
    date.fromisoformat(as_of)
    issuers = []
    sources = load_sources(data_dir)
    for source, document in sources:
        rows = [select_annual(document, source, year, as_of) for year in YEARS]
        for index, row in enumerate(rows):
            row["ratios"] = calculate(row, rows[index - 1] if index else None)
            v = row["values"]
            row["quality"] = {"missing_fields": [k for k, value in v.items() if value is None],
                              "balance_sheet_residual_usd": difference(v["total_assets"],
                                  v["total_liabilities"] + v["consolidated_equity"] if v["total_liabilities"] is not None and v["consolidated_equity"] is not None else None)}
        issuers.append({"ticker": source["ticker"], "name": document["entityName"],
                        "context": CONTEXT[source["ticker"]], "years": rows, "monitoring": monitoring(rows)})
    peers = {}
    for key, (label, formula, unit) in RATIOS.items():
        observations = [issuer["years"][-1]["ratios"][key] for issuer in issuers if issuer["years"][-1]["ratios"][key] is not None]
        peers[key] = {"label": label, "formula": formula, "unit": unit,
                      "median": median(observations) if observations else None, "n": len(observations)}
    for issuer in issuers:
        issuer["review"] = review_text(issuer, peers, as_of)
    return {"case_id": "us-beverages-fy2022-2024", "title": "Public filings credit review: beverages",
            "as_of": as_of, "years": list(YEARS), "currency": "USD", "unit": "USD (unscaled)",
            "data_status": "Historical reported financials; current raw cache filtered to original annual filings at cutoff",
            "selection_policy": "Earliest 10-K annual revenue accession for each period ending 2022–2024, filed by cutoff; exact accession, USD unit and start/end for other facts. Later amendments and restatements are excluded.",
            "limitations": ["Three US-listed corporations; no Asia, financial institution or fund case is evidenced here.",
                            "Small selected peer set, including a foods business and differing fiscal calendars; medians are not sector benchmarks.",
                            "Research screens and narrative are generated analysis. No agency rating, calibrated PD, lending limit, trade, client exposure or onboarding approval is claimed.",
                            "Standard US-GAAP facts only. Missing concepts remain unavailable. Liabilities may be derived using equity including noncontrolling interests.",
                            "FCF is defined here as operating cash flow less cash capital expenditure; it may differ from issuer non-GAAP measures. No one-off adjustments are applied.",
                            "Annual historical monitoring only; current news, legal agreements, collateral and counterparty limits are outside this dataset.",
                            "A cash ratio below one or a screening alert is not proof of default or a credit recommendation."],
            "sources": [source for source, _ in sources], "issuers": issuers, "peer_metrics": peers}


def display(value: float | None, unit: str = "%") -> str:
    if value is None:
        return "unavailable"
    return f"{value * 100:.2f}%" if unit == "%" else f"{value:.2f}x"


def review_text(issuer: dict, peers: dict, as_of: str) -> str:
    latest = issuer["years"][-1]
    v, r = latest["values"], latest["ratios"]
    lines = [f"# {issuer['name']} ({issuer['ticker']}) — historical credit review",
             f"\nInformation cutoff: {as_of}. Fiscal period: {latest['period_start']} to {latest['period_end']}. All amounts USD.",
             "\n## Business and comparability", issuer["context"]["profile"],
             "\n## Reported financial performance",
             f"Revenue: ${v['revenue']:,.0f}; annual growth: {display(r['revenue_growth'])}. Operating margin: {display(r['operating_margin'])}; selected peer median: {display(peers['operating_margin']['median'])} (n={peers['operating_margin']['n']}, including this issuer).",
             f"Current ratio: {display(r['current_ratio'], 'x')}; liabilities/assets: {display(r['liabilities_to_assets'])}. Operating cash flow margin: {display(r['cfo_margin'])}; free cash flow margin: {display(r['fcf_margin'])}.",
             "\n## Filing context", issuer["context"]["event"],
             f"Source: {issuer['context']['source']} ({issuer['context']['section']}).",
             "\n## Monitoring and follow-up"]
    for check in issuer["monitoring"]:
        lines.append(f"- {check['status']}: {check['rule']}. Illustrative research screen, not a covenant test.")
    lines.extend(["\nReview liquidity facilities, debt maturity disclosures, recurring cash generation and the cited exceptional items before reaching a lending decision.",
                  "\n## Evidence and decision limits", f"Annual filing: {latest['filing_url']}; filed {latest['filed']}.",
                  "Each extracted financial field retains its concept, accession, period, unit and source checksum. This generated review is an analytical draft; no credit rating, calibrated PD or facility approval is assigned."])
    return "\n\n".join(lines) + "\n"


def statement_csv(issuer: dict) -> str:
    output = io.StringIO()
    writer = csv.DictWriter(output, fieldnames=["period", *FIELDS, "net_interest_and_other", "total_equity"])
    writer.writeheader()
    for row in issuer["years"]:
        writer.writerow({"period": str(row["year"]), **row["values"], "total_equity": row["values"]["equity"]})
    return output.getvalue()


def export_case(case: dict, destination: Path) -> None:
    destination.mkdir(parents=True, exist_ok=True)
    (destination / "case_study.json").write_text(json.dumps(case, indent=2, allow_nan=False) + "\n")
    with (destination / "ratios.csv").open("w", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=["ticker", "year", "period_end", *RATIOS])
        writer.writeheader()
        for issuer in case["issuers"]:
            for row in issuer["years"]:
                writer.writerow({k: row[k] for k in ("ticker", "year", "period_end")} | row["ratios"])
    for issuer in case["issuers"]:
        (destination / f"{issuer['ticker']}_statements.csv").write_text(statement_csv(issuer))
        (destination / f"{issuer['ticker']}_credit_review.md").write_text(issuer["review"])
    # The workbook contains reproducible Excel formulas referencing reported inputs.
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Financials"
    financial_fields = [*FIELDS, "net_interest_and_other"]
    fields = ["ticker", "year", *financial_fields]
    sheet.append(fields)
    for issuer in case["issuers"]:
        for row in issuer["years"]:
            sheet.append([row["ticker"], row["year"], *[row["values"].get(key) for key in financial_fields]])
    ratios = workbook.create_sheet("Ratios")
    ratios.append(["ticker", "year", "operating_margin", "current_ratio", "liabilities_to_assets", "cfo_margin", "fcf_margin"])
    col = {key: get_column_letter(fields.index(key) + 1) for key in FIELDS}
    def excel_ratio(index: int, numerator: str, denominator: str, subtract: str | None = None) -> str:
        def ref(key: str) -> str:
            return f"Financials!{col[key]}{index}"
        required = [numerator, denominator] + ([subtract] if subtract else [])
        checks = ",".join(f"ISNUMBER({ref(key)})" for key in required)
        expression = f"({ref(numerator)}-{ref(subtract)})" if subtract else ref(numerator)
        return f'=IF(AND({checks},{ref(denominator)}>0),{expression}/{ref(denominator)},"")'

    for index in range(2, sheet.max_row + 1):
        ratios.append([f"=Financials!A{index}", f"=Financials!B{index}",
                       excel_ratio(index, "operating_income", "revenue"),
                       excel_ratio(index, "current_assets", "current_liabilities"),
                       excel_ratio(index, "total_liabilities", "total_assets"),
                       excel_ratio(index, "operating_cash_flow", "revenue"),
                       excel_ratio(index, "operating_cash_flow", "revenue", "capital_expenditure")])
    evidence = workbook.create_sheet("Evidence")
    evidence.append(["ticker", "year", "field", "kind", "concept_or_formula", "accession", "filed", "source_sha256"])
    for issuer in case["issuers"]:
        for row in issuer["years"]:
            for field, fact in row["evidence"].items():
                evidence.append([issuer["ticker"], row["year"], field, fact["kind"], fact.get("concept", fact.get("formula", fact.get("reason"))), fact.get("accn"), fact.get("filed"), fact.get("source_sha256")])
    notes = workbook.create_sheet("Read me")
    for text in [case["title"], f"Information cutoff: {case['as_of']}; all financials USD, ratios are fractions", case["selection_policy"], *case["limitations"]]:
        notes.append([text])
    for page in workbook:
        page.freeze_panes = "C2"
        page.auto_filter.ref = page.dimensions
        for column in page.columns:
            page.column_dimensions[column[0].column_letter].width = min(55, max(16, len(str(column[0].value)) + 2))
    workbook.save(destination / "credit_case.xlsx")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--refresh", action="store_true", help="Download SEC sources; default rebuilds offline")
    parser.add_argument("--user-agent", default="", help="Project and real contact email, required for refresh")
    parser.add_argument("--data-dir", type=Path, default=DATA)
    parser.add_argument("--output", type=Path, default=DATA / "outputs")
    args = parser.parse_args()
    if args.refresh:
        download_sources(args.data_dir, args.user_agent)
    case = build_case(args.data_dir)
    export_case(case, args.output)
    rows = [row for issuer in case["issuers"] for row in issuer["years"]]
    result = {"issuers": len(case["issuers"]), "annual_rows": len(rows),
              "reported_facts": sum(f["kind"] == "reported" for row in rows for f in row["evidence"].values()),
              "derived_facts": sum(f["kind"] == "derived" for row in rows for f in row["evidence"].values()),
              "missing_facts": sum(f["kind"] == "unavailable" for row in rows for f in row["evidence"].values()),
              "screen_alerts": sum(c["status"] == "review" for i in case["issuers"] for c in i["monitoring"]),
              "as_of": case["as_of"], "output": str(args.output)}
    print(json.dumps(result, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
