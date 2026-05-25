import copy
import gzip
import io
import json
import zipfile

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
import pytest

from sentinel_risk.public_filings import (
    AS_OF, DATA, build_case, calculate, divide, export_case, load_sources, select_annual, statement_csv,
)
from sentinel_risk.io import load_statement
from app.routers.public_filings import router


@pytest.fixture(scope="module")
def case():
    return build_case()


def issuer(case, ticker):
    return next(item for item in case["issuers"] if item["ticker"] == ticker)


def test_real_filings_agree_with_published_2024_statements(case):
    expected = {"KO": (47061000000, 9992000000, 6805000000, 2064000000),
                "PEP": (91854000000, 12887000000, 12507000000, 5318000000),
                "KDP": (15351000000, 2591000000, 2219000000, 563000000)}
    for ticker, figures in expected.items():
        row = issuer(case, ticker)["years"][-1]
        assert tuple(row["values"][key] for key in ("revenue", "operating_income", "operating_cash_flow", "capital_expenditure")) == figures
        assert row["filed"] <= AS_OF


def test_period_accession_and_cashflow_integrity(case):
    for company in case["issuers"]:
        for row in company["years"]:
            assert row["quality"]["balance_sheet_residual_usd"] == 0
            for field, fact in row["evidence"].items():
                if fact["kind"] == "reported":
                    assert fact["accn"] == row["accession"]
                    assert fact["end"] == row["period_end"]
                    assert fact["unit"] == "USD"
                    if fact["start"]:
                        assert fact["start"] == row["period_start"]
                    assert row["values"][field] == fact["val"]
    assert issuer(case, "PEP")["years"][-1]["period_end"] == "2024-12-28"


def test_liabilities_use_consolidated_equity_including_minority(case):
    row = issuer(case, "KO")["years"][-1]
    assert row["values"]["total_liabilities"] == 74177000000
    assert row["evidence"]["total_liabilities"]["kind"] == "derived"
    assert row["values"]["total_liabilities"] != row["values"]["total_assets"] - row["values"]["equity"]


def test_missing_concept_stays_missing(case):
    row = issuer(case, "KDP")["years"][-1]
    assert row["values"]["interest_expense"] is None
    assert row["evidence"]["interest_expense"]["kind"] == "unavailable"
    assert divide(1, 0) is None
    assert divide(1, -1) is None
    assert divide(None, 1) is None


def test_pepsico_net_interest_is_not_exported_as_gross_interest(case, tmp_path):
    pep = issuer(case, "PEP")
    assert pep["years"][-1]["values"]["net_interest_and_other"] == 919000000
    for row in pep["years"]:
        assert row["values"]["interest_expense"] is None
        assert row["evidence"]["net_interest_and_other"]["issuer_label"] == "Net interest expense and other"
    source = tmp_path / "PEP_statements.csv"
    source.write_text(statement_csv(pep))
    # A downstream strict coverage calculator must reject the missing gross
    # input, rather than calculating a false 14.02x coverage ratio from 919m.
    with pytest.raises(ValueError, match="non-numeric or empty"):
        load_statement(source)


def test_ratio_and_peer_math_with_independent_inputs(case):
    row = issuer(case, "KO")["years"][-1]
    assert row["ratios"]["operating_margin"] == pytest.approx(9992 / 47061)
    assert row["ratios"]["fcf_margin"] == pytest.approx((6805 - 2064) / 47061)
    assert row["ratios"]["revenue_growth"] == pytest.approx(47061 / 45754 - 1)
    # KDP is the middle operating margin in this three-company peer set.
    assert case["peer_metrics"]["operating_margin"]["median"] == pytest.approx(2591 / 15351)
    assert case["peer_metrics"]["operating_margin"]["n"] == 3
    assert issuer(case, "KO")["years"][0]["ratios"]["revenue_growth"] is None


def test_missing_cashflow_never_creates_fake_fcf(case):
    row = copy.deepcopy(issuer(case, "KO")["years"][-1])
    row["values"]["capital_expenditure"] = None
    assert calculate(row)["fcf_margin"] is None


def test_monitoring_flags_are_calculated_and_not_covenants(case):
    assert sum(check["status"] == "review" for company in case["issuers"] for check in company["monitoring"]) == 5
    cfo = next(check for check in issuer(case, "KO")["monitoring"] if check["area"] == "Cash flow")
    assert cfo["value"] == pytest.approx(6805 / 11599 - 1)
    assert "not a covenant" in cfo["basis"]


def test_selector_excludes_quarters_later_restatements_and_other_units():
    source = {"ticker": "TEST", "cik": "0000000001", "sha256": "test"}
    annual = {"start": "2024-01-01", "end": "2024-12-31", "val": 100, "accn": "first", "form": "10-K", "filed": "2025-02-01", "fy": 2025}
    quarter = annual | {"start": "2024-10-01", "val": 30, "filed": "2025-01-01", "accn": "quarter"}
    restatement = annual | {"val": 120, "filed": "2025-03-01", "accn": "later"}
    future = annual | {"val": 200, "filed": "2026-02-01", "accn": "future"}
    doc = {"facts": {"us-gaap": {"Revenues": {"units": {"USD": [future, quarter, restatement, annual], "EUR": [annual | {"val": 999}]}}}}}
    row = select_annual(doc, source, 2024)
    assert row["accession"] == "first"
    assert row["values"]["revenue"] == 100  # fy=2025 must not exclude a 2024 period.
    with pytest.raises(ValueError, match="No annual revenue"):
        select_annual(doc, source, 2024, "2024-12-31")


def test_selector_rejects_conflicting_same_period_values():
    source = {"ticker": "TEST", "cik": "0000000001", "sha256": "test"}
    annual = {"start": "2024-01-01", "end": "2024-12-31", "val": 100, "accn": "first", "form": "10-K", "filed": "2025-02-01"}
    doc = {"facts": {"us-gaap": {"Revenues": {"units": {"USD": [annual, annual | {"val": 101}]}}}}}
    with pytest.raises(ValueError, match="Ambiguous"):
        select_annual(doc, source, 2024)


def test_corrupt_source_is_rejected(tmp_path):
    sources = json.loads((DATA / "sources.json").read_text())
    sources[0]["cache"] = "bad.gz"
    (tmp_path / "sources.json").write_text(json.dumps(sources))
    (tmp_path / "bad.gz").write_bytes(gzip.compress(b'{"cik":21344}'))
    with pytest.raises(ValueError, match="checksum mismatch"):
        load_sources(tmp_path)


@pytest.mark.parametrize("change", ["relabel", "duplicate", "missing"])
def test_manifest_requires_correct_unique_complete_issuers(tmp_path, change):
    sources = json.loads((DATA / "sources.json").read_text())
    if change == "relabel":
        sources[0] = sources[1] | {"ticker": "KO"}
    elif change == "duplicate":
        sources[0] = sources[1]
    else:
        sources.pop()
    (tmp_path / "sources.json").write_text(json.dumps(sources))
    with pytest.raises(ValueError, match="ticker/CIK mismatch|each configured issuer"):
        load_sources(tmp_path)


@pytest.mark.parametrize("path_kind", ["absolute", "traversal", "symlink"])
def test_cache_paths_cannot_escape_dataset(tmp_path, path_kind):
    directory = tmp_path / "case"
    directory.mkdir()
    outside = tmp_path / "outside.gz"
    outside.write_bytes(b"outside")
    sources = json.loads((DATA / "sources.json").read_text())
    if path_kind == "absolute":
        sources[0]["cache"] = str(outside)
    elif path_kind == "traversal":
        sources[0]["cache"] = "../outside.gz"
    else:
        (directory / "link.gz").symlink_to(outside)
        sources[0]["cache"] = "link.gz"
    (directory / "sources.json").write_text(json.dumps(sources))
    with pytest.raises(ValueError, match="Cache path escapes dataset"):
        load_sources(directory)


def test_excel_ratio_formulas_guard_all_inputs_and_denominators(case, tmp_path):
    modified = copy.deepcopy(case)
    modified["issuers"][0]["years"][0]["values"]["capital_expenditure"] = None
    modified["issuers"][0]["years"][1]["values"]["revenue"] = -1
    export_case(modified, tmp_path)
    book = load_workbook(tmp_path / "credit_case.xlsx", data_only=False)
    financials = book["Financials"]
    columns = {cell.value: cell.column_letter for cell in financials[1]}
    assert financials[f'{columns["capital_expenditure"]}2'].value is None
    for row_number in range(2, 11):
        for cell in book["Ratios"][row_number][2:]:
            formula = cell.value
            # Every referenced cell, including both terms of the FCF numerator,
            # must appear in a numeric guard; denominator must be positive.
            import re
            references = set(re.findall(r"Financials![A-Z]+\d+", formula))
            assert all(f"ISNUMBER({reference})" in formula for reference in references)
            denominator = formula.rsplit("/", 1)[1].split(",", 1)[0]
            assert f"{denominator}>0" in formula
            assert formula.endswith(',"")')


def test_case_is_reproducible(case):
    assert build_case() == case


def test_api_returns_case_and_downloadable_calculation_pack(case):
    app = FastAPI()
    app.include_router(router, prefix="/public-filings")
    client = TestClient(app)
    response = client.get("/public-filings/case")
    assert response.status_code == 200
    assert response.json() == case
    assert client.get("/public-filings/UNKNOWN/statements").status_code == 404
    csv_response = client.get("/public-filings/KO/statements")
    assert csv_response.status_code == 200
    assert len(csv_response.text.strip().splitlines()) == 4
    download = client.get("/public-filings/export")
    assert download.status_code == 200
    archive = zipfile.ZipFile(io.BytesIO(download.content))
    assert len(archive.namelist()) == 10
    assert json.loads(archive.read("case_study.json")) == case
    workbook = load_workbook(io.BytesIO(archive.read("credit_case.xlsx")), data_only=False)
    assert workbook["Financials"].max_row == 10
    assert workbook["Evidence"].max_row == 157
    assert workbook["Ratios"]["C2"].value.startswith("=IF(AND(")
    assert "no credit rating, calibrated PD or facility approval" in case["issuers"][0]["review"]
